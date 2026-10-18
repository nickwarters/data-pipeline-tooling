#!/usr/bin/env python3
"""Build a Question Bank artifact from one sheet of an Excel workbook.

    python3 scripts/import_bank_from_excel.py --file questions.xlsx \
        --sheet Questions --slug complaints

Writes `case-types/banks/{slug}.txt` from the sheet's rows, then publishes it
with `node scripts/publish-bank.js {slug}`: that mints the new version, writes
its immutable `{slug}.{version}.txt` copy and appends the version to the bank's
`history`. Publishing is delegated rather than reimplemented here so there is
one rule for what a version is. Pass `--no-publish` to write the bank only.

The first non-empty row is the header row. Headers are matched ignoring case,
spaces and punctuation; `text` (or `Question`) is the only required column:

| Column               | Cell                                                    |
| -------------------- | ------------------------------------------------------- |
| `id`                 | blank: reuse the id of an existing question with the    |
|                      | same text, otherwise mint the next `q-...-NNNN`         |
| `text` / `Question`  | the question wording                                    |
| `questionGroup`      | the group it renders under                              |
| `category`           |                                                         |
| `responseType`       | `outcome` (blank default), `yes-no-na`,                 |
|                      | `single-choice` or `multi-choice`                       |
| `options`            | one per line (or `|`-separated); choice types only      |
| `optionOutcomes`     | `Option = outcome-id`, one per line                     |
| `labels`             | label names or ids, one per line                        |
| `remediationActions` | one per line                                            |
| `showWhen`           | the `showWhen` object as JSON                           |
| `deprecated`         | yes/true/1 to retire the question                       |

An existing bank is the starting point, not something to replace wholesale:
its label, Labels and Outcome vocabulary are kept, a question keeps any field
the sheet has no column for, and a question the sheet no longer lists is kept
as `deprecated` because Question Definitions are never deleted — Cases already
answered against it must keep resolving its id.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

SCRIPT_ROOT = Path(__file__).resolve().parents[1]

RESPONSE_TYPES = ("yes-no-na", "single-choice", "multi-choice", "outcome")
CHOICE_TYPES = ("single-choice", "multi-choice")

# Normalised header -> bank field.
COLUMNS = {
    "id": "id",
    "questionid": "id",
    "text": "text",
    "question": "text",
    "questiontext": "text",
    "questiongroup": "questionGroup",
    "group": "questionGroup",
    "category": "category",
    "responsetype": "responseType",
    "type": "responseType",
    "options": "options",
    "responseoptions": "options",
    "optionoutcomes": "optionOutcomes",
    "labels": "labelIds",
    "labelids": "labelIds",
    "remediationactions": "remediationActions",
    "remediation": "remediationActions",
    "showwhen": "showWhen",
    "deprecated": "deprecated",
}

# The order the Question Bank editor writes a question's keys in.
QUESTION_KEY_ORDER = (
    "id",
    "text",
    "category",
    "questionGroup",
    "labelIds",
    "responseType",
    "options",
    "optionOutcomes",
    "showWhen",
    "remediationActions",
    "disallowFreeFormRemediation",
    "deprecated",
)


class SheetError(Exception):
    """A problem with the workbook the operator has to fix; reported, not traced."""


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build case-types/banks/{slug}.txt from an Excel sheet and publish it."
    )
    parser.add_argument("--file", required=True, type=Path, help="the .xlsx workbook")
    parser.add_argument("--sheet", required=True, help="the sheet holding the questions")
    parser.add_argument("--slug", required=True, help="the Case Type slug, e.g. complaints")
    parser.add_argument(
        "--label", help="the bank's display label (required when the bank is new)"
    )
    parser.add_argument(
        "--no-publish",
        action="store_true",
        help="write the bank only; do not mint a version or advance its history",
    )
    parser.add_argument(
        "--root",
        type=Path,
        default=SCRIPT_ROOT,
        help="the platform_frontend directory (default: the one this script is in)",
    )
    return parser.parse_args(argv)


def normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def cell_text(value: Any) -> str:
    """A cell as text: Excel hands back ints as floats and leaves stray spaces."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def split_lines(value: Any) -> list[str]:
    return [part.strip() for part in re.split(r"[\n|]", cell_text(value)) if part.strip()]


def read_sheet(path: Path, sheet: str) -> list[dict[str, Any]]:
    """The sheet's rows as `{field: raw cell}`, keyed by the bank field each header names."""
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise SheetError(
            "openpyxl is not installed — activate the repository venv "
            "(`pip install -r requirements-dev.txt` at the repository root)"
        ) from error

    if not path.is_file():
        raise SheetError(f"workbook not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise SheetError(
                f"sheet {sheet!r} not found in {path.name} — it has: "
                + ", ".join(repr(name) for name in workbook.sheetnames)
            )
        rows = [
            row
            for row in workbook[sheet].iter_rows(values_only=True)
            if any(cell_text(cell) for cell in row)
        ]
    finally:
        workbook.close()

    if not rows:
        raise SheetError(f"sheet {sheet!r} is empty")
    header, *body = rows

    fields: list[str | None] = []
    for raw in header:
        name = normalise_header(raw)
        if not name:
            fields.append(None)
            continue
        if name not in COLUMNS:
            raise SheetError(
                f"column {raw!r} is not one this importer understands — use one of: "
                + ", ".join(sorted(set(COLUMNS.values())))
            )
        if COLUMNS[name] in fields:
            raise SheetError(f"two columns both name {COLUMNS[name]!r}")
        fields.append(COLUMNS[name])
    if "text" not in fields:
        raise SheetError("the sheet has no question text column (`text` or `Question`)")

    return [
        {field: cell for field, cell in zip(fields, row) if field is not None}
        for row in body
    ]


def truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return cell_text(value).lower() in ("y", "yes", "true", "1", "x")


class IdMinter:
    """Mints the next id in the bank's existing `prefix` + zero-padded number style."""

    def __init__(self, slug: str, existing_ids: list[str]) -> None:
        numbered = [m for m in (re.fullmatch(r"(.*?)(\d+)", i) for i in existing_ids) if m]
        prefixes = [m.group(1) for m in numbered]
        self.prefix = max(set(prefixes), key=prefixes.count) if prefixes else f"q-{slug}-"
        same = [m for m in numbered if m.group(1) == self.prefix]
        self.width = max((len(m.group(2)) for m in same), default=4)
        self.next = max((int(m.group(2)) for m in same), default=0) + 1
        self.taken = set(existing_ids)

    def mint(self) -> str:
        while True:
            candidate = f"{self.prefix}{self.next:0{self.width}d}"
            self.next += 1
            if candidate not in self.taken:
                self.taken.add(candidate)
                return candidate


def resolve_labels(cell: Any, labels: list[dict[str, Any]], where: str) -> list[str]:
    by_key = {}
    for label in labels:
        by_key[label["id"]] = label["id"]
        by_key[str(label.get("name", "")).lower()] = label["id"]
    ids = []
    for name in split_lines(cell):
        label_id = by_key.get(name) or by_key.get(name.lower())
        if not label_id:
            raise SheetError(
                f"{where}: label {name!r} is not in the bank — it has: "
                + ", ".join(repr(label.get("name")) for label in labels)
                + " (add new Labels in the Question Bank editor first)"
            )
        ids.append(label_id)
    return ids


def parse_option_outcomes(cell: Any, outcome_ids: set[str], where: str) -> dict[str, str]:
    mapping = {}
    for line in split_lines(cell):
        option, sep, outcome = line.rpartition("=")
        if not sep or not option.strip() or not outcome.strip():
            raise SheetError(f"{where}: optionOutcomes line {line!r} is not `Option = outcome-id`")
        if outcome_ids and outcome.strip() not in outcome_ids:
            raise SheetError(
                f"{where}: {outcome.strip()!r} is not one of the bank's Outcomes: "
                + ", ".join(sorted(outcome_ids))
            )
        mapping[option.strip()] = outcome.strip()
    return mapping


def build_question(
    qid: str, row: dict[str, Any], base: dict[str, Any], bank: dict[str, Any], where: str
) -> dict[str, Any]:
    """One Question Definition: the sheet's columns over the existing question's fields."""
    question = {**base, "id": qid}

    def set_or_clear(field: str, value: Any) -> None:
        if value in (None, "", [], {}):
            question.pop(field, None)
        else:
            question[field] = value

    question["text"] = cell_text(row["text"])
    if not question["text"]:
        raise SheetError(f"{where}: has no question text")
    for field in ("questionGroup", "category"):
        if field in row:
            set_or_clear(field, cell_text(row[field]))

    if cell_text(row.get("responseType")):
        question["responseType"] = cell_text(row["responseType"]).lower()
    question.setdefault("responseType", "outcome")
    response_type = question["responseType"]
    if response_type not in RESPONSE_TYPES:
        raise SheetError(
            f"{where}: responseType {response_type!r} — use one of: {', '.join(RESPONSE_TYPES)}"
        )

    if base.get("responseType") == "outcome" and response_type != "outcome":
        # Those options were derived from the Outcomes, not authored for this type.
        question.pop("options", None)
        question.pop("optionOutcomes", None)
    if "options" in row:
        set_or_clear("options", split_lines(row["options"]))
    if "optionOutcomes" in row:
        outcome_ids = {o["id"] for o in bank.get("outcomeOptions", [])}
        set_or_clear("optionOutcomes", parse_option_outcomes(row["optionOutcomes"], outcome_ids, where))
    if response_type == "outcome":
        # An Outcome question answers in the Case Type's Outcomes; its options are
        # derived from them, exactly as the Question Bank editor compiles them.
        # An option is the Outcome's id, which is what an Answer stores, so a
        # reworded Outcome leaves every recorded Answer pointing at it.
        if split_lines(row.get("options")) or cell_text(row.get("optionOutcomes")):
            raise SheetError(
                f"{where}: an `outcome` question takes its options from the bank's Outcomes "
                "— leave options/optionOutcomes blank"
            )
        outcomes = bank.get("outcomeOptions", [])
        set_or_clear("options", [o["id"] for o in outcomes])
        set_or_clear("optionOutcomes", {o["id"]: o["id"] for o in outcomes})
    elif response_type in CHOICE_TYPES and not question.get("options"):
        raise SheetError(f"{where}: a {response_type} question needs at least one option")

    if "labelIds" in row:
        set_or_clear("labelIds", resolve_labels(row["labelIds"], bank.get("labels", []), where))
    if "remediationActions" in row:
        set_or_clear("remediationActions", split_lines(row["remediationActions"]))
    if "showWhen" in row:
        text = cell_text(row["showWhen"])
        try:
            show_when = json.loads(text) if text else None
        except json.JSONDecodeError as error:
            raise SheetError(f"{where}: showWhen is not valid JSON — {error}") from error
        if show_when is not None and not isinstance(show_when, dict):
            raise SheetError(f"{where}: showWhen must be a JSON object")
        set_or_clear("showWhen", show_when)
    question["deprecated"] = truthy(row["deprecated"]) if "deprecated" in row else question.get("deprecated", False)

    ordered = {key: question[key] for key in QUESTION_KEY_ORDER if key in question}
    return {**ordered, **{k: v for k, v in question.items() if k not in ordered}}


def build_bank(
    rows: list[dict[str, Any]], existing: dict[str, Any] | None, slug: str, label: str | None
) -> tuple[dict[str, Any], list[str]]:
    """The new bank, and the ids of existing questions it retired for not being listed."""
    existing = existing or {}
    bank_label = label or existing.get("label")
    if not bank_label:
        raise SheetError(f"no bank exists for {slug!r} yet — pass --label to name it")
    context = {
        "labels": existing.get("labels", []),
        "outcomeOptions": existing.get("outcomeOptions", []),
    }

    old = {q["id"]: q for q in existing.get("questions", [])}
    minter = IdMinter(slug, list(old))
    unclaimed_by_text = {}
    for q in existing.get("questions", []):
        unclaimed_by_text.setdefault(q.get("text", "").strip(), q["id"])

    # Reserve every id the sheet spells out, so a minted one never collides with a later row.
    minter.taken.update(cell_text(row.get("id")) for row in rows)
    questions = []
    seen: set[str] = set()
    for number, row in enumerate(rows, start=2):
        where = f"row {number}"
        qid = cell_text(row.get("id"))
        if not qid:
            qid = unclaimed_by_text.get(cell_text(row["text"])) or ""
            if qid in seen:
                qid = ""
        qid = qid or minter.mint()
        if qid in seen:
            raise SheetError(f"{where}: id {qid!r} appears more than once")
        seen.add(qid)
        questions.append(build_question(qid, row, old.get(qid, {}), context, f"{where} ({qid})"))

    retired = [qid for qid in old if qid not in seen]
    questions += [{**old[qid], "deprecated": True} for qid in retired]

    bank: dict[str, Any] = {"slug": slug, "label": bank_label}
    for key in ("version", "history"):
        if key in existing:
            bank[key] = existing[key]
    bank["questions"] = questions
    for key in ("labels", "outcomeOptions", "defaultOutcomeId"):
        if existing.get(key):
            bank[key] = existing[key]
    return bank, retired


def publish(root: Path, slug: str) -> None:
    node = shutil.which("node")
    if not node:
        raise SheetError(
            f"node is not on PATH, so the bank was written but not published — "
            f"run `node scripts/publish-bank.js {slug}` from {root}"
        )
    # Resolved, because publish-bank.js only runs when its argv path is its own real
    # path; through a symlink (macOS's /tmp, /var) it exits 0 having done nothing.
    script = (root / "scripts" / "publish-bank.js").resolve()
    result = subprocess.run([node, str(script), slug], cwd=root)
    if result.returncode != 0:
        raise SheetError(f"publish-bank.js failed for {slug!r} (exit {result.returncode})")
    bank = json.loads((root / "case-types" / "banks" / f"{slug}.txt").read_text(encoding="utf-8"))
    if not bank.get("version") or bank.get("history", [{}])[-1].get("version") != bank["version"]:
        raise SheetError(f"publish-bank.js ran but {slug!r} declares no published version")


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]*", args.slug):
        print(f"error: slug {args.slug!r} must be lower-case letters, digits and -", file=sys.stderr)
        return 2
    bank_path = args.root / "case-types" / "banks" / f"{args.slug}.txt"
    try:
        existing = json.loads(bank_path.read_text(encoding="utf-8")) if bank_path.exists() else None
        bank, retired = build_bank(read_sheet(args.file, args.sheet), existing, args.slug, args.label)
        bank_path.parent.mkdir(parents=True, exist_ok=True)
        bank_path.write_text(
            json.dumps(bank, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n"
        )
        active = sum(1 for q in bank["questions"] if not q["deprecated"])
        print(f"wrote {bank_path}: {len(bank['questions'])} questions, {active} active")
        if retired:
            print(f"  kept as deprecated (not in the sheet): {', '.join(retired)}")
        if not bank.get("outcomeOptions"):
            print(
                "  note: the bank has no Outcome vocabulary yet — add outcomeOptions and "
                "defaultOutcomeId before the Case Type will load"
            )
        if not args.no_publish:
            sys.stdout.flush()  # keep our lines ahead of node's in the terminal
            publish(args.root, args.slug)
    except SheetError as error:
        print(f"error: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
